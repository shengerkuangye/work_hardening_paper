#!/usr/bin/env python3
"""Targeted replacement of figure 12 and its discussion in the existing Word.
Backup and QA artifacts stay in .codex_tmp; no second manuscript is published.
"""
from pathlib import Path
from copy import deepcopy
from datetime import datetime
import hashlib
import json
import shutil
import tempfile
from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / '2026-09-27/2026-09-27-ta4.docx'
OUT = ROOT / 'results/tem_same_field_2026-10-01'
TEXT = {
    107: '图 12 展示了旋锻态 Y-5.00 在三组衍射条件下的明场 STEM 图像及配套电子衍射花样。低倍图中共同的孔缘轮廓及定位特征表明，三组观察视场具有重叠区域（图 12a–c），黄色框分别标示对应高倍图像的覆盖范围。高倍图中可见弯曲、交织的位错线及其局部聚集（图 12d–f）；不同衍射条件下，线状衬度的清晰程度及其与带状背景的叠加方式有所不同。配套衍射花样的主要系统排按 α-Ti 晶体学关系初步归属于（0002）、{10−10}和{10−11}反射（图 12g–i）。上述结果展示了旋锻态复杂的位错组态及其衍射衬度特征，与图 10 所示局部缠结形貌相一致。',
    109: '图 12  旋锻态 Y-5.00 重叠视场的明场 STEM 图像与配套电子衍射花样。（a–c）低倍图像，编号为 0026、0029、0032；（d–f）各列对应的高倍图像，编号为 0028、0031、0034；（g–i）配套衍射花样，编号为 0019、0020、0022。青色 L 标示共同孔缘定位特征，黄色框标示高倍视场范围。比例尺为 1 μm（a–c）、500 nm（d–f）和 10 nm⁻¹（g–i）。星号表示系统排指数为初步标定；四指数中的负号表示相应指数上方的横线。',
    110: '位错类型的分析依据不同衍射条件下的衬度响应及 g·b 消光关系展开，其中 g 为衍射矢量，b 为柏氏矢量。（0002）反射与基面内反射的联合分析可分别约束柏氏矢量的 c 轴分量和基面内分量。在〈a〉、〈c〉和〈c+a〉型位错的候选范围内，同一位错在两类适当反射下均保持衬度，可支持其同时具有 a 与 c 分量的解释。结合图 8、9 所示滑移几何条件，含 c 分量位错参与局部应变协调可作为旋锻变形的机制假设加以讨论。对于〈c+a〉型位错，其 c 轴分量可补充〈a〉型滑移的变形协调能力[6]。',
}


def main():
    doc = Document(TARGET)
    assert doc.paragraphs[107].text.startswith('图 12')
    assert doc.paragraphs[109].text.startswith('图 12')
    scratch = Path(tempfile.mkdtemp(prefix='tem_same_field_', dir=ROOT/'.codex_tmp'))
    shutil.copy2(TARGET, scratch/'before.docx')
    before = hashlib.sha256(TARGET.read_bytes()).hexdigest()
    changes=[]
    for idx, text in TEXT.items():
        p=doc.paragraphs[idx]
        changes.append(dict(paragraph=idx, before=p.text, after=text))
        rpr=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
        p.clear()
        run=p.add_run(text)
        if rpr is not None:
            run._r.insert(0,rpr)
    p=doc.paragraphs[108]
    p.clear()
    p.add_run().add_picture(str(OUT/'fig12_y5_same_field_diffraction.png'), width=Inches(6.25))
    p.paragraph_format.keep_with_next=True
    # The existing manuscript also has an obsolete TEM panel embedded inside
    # the first conclusion paragraph. Remove only that drawing, not its text.
    stale = doc.paragraphs[123]
    removed = []
    for drawing in list(stale._p.xpath('.//w:drawing')):
        for blip in drawing.xpath('.//a:blip'):
            rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            removed.append(str(doc.part.related_parts[rid].partname))
        drawing.getparent().remove(drawing)
    doc.save(scratch/'edited.docx')
    assert hashlib.sha256(TARGET.read_bytes()).hexdigest()==before, 'Document changed during edit'
    # Preserve every untargeted paragraph and table at the XML level.
    original=Document(scratch/'before.docx')
    check=Document(scratch/'edited.docx')
    assert len(original.paragraphs)==len(check.paragraphs)
    for idx,(a,b) in enumerate(zip(original.paragraphs,check.paragraphs)):
        if idx not in [107,108,109,110,123]:
            assert a._p.xml==b._p.xml, f'Unexpected paragraph change {idx}'
    assert original.paragraphs[123].text==check.paragraphs[123].text
    assert [t._tbl.xml for t in original.tables]==[t._tbl.xml for t in check.tables]
    shutil.copy2(scratch/'edited.docx',TARGET)
    (scratch/'changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2))
    (OUT/'manuscript_text.txt').write_text('\n\n'.join(TEXT.values())+'\n')
    (OUT/'word_update_record.json').write_text(json.dumps(dict(target=str(TARGET),backup=str(scratch/'before.docx'),changed_paragraphs=list(TEXT),figure_paragraph=108,obsolete_duplicate_tem_drawing_removed=removed,conclusion_text_preserved=True,untargeted_xml_verified=True),ensure_ascii=False,indent=2)+'\n')
    print(scratch)


if __name__=='__main__':
    main()
